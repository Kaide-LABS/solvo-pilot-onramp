"""Stage 2 — Excel extractor. Implements PHASE_2_SPEC.md §6.2.

Reads .xlsx via openpyxl (read_only, data_only), produces a structured cell
list (NOT the raw blob), and asks gemini-3.1-flash-lite to return a
NormalizedRatesheet that conforms to the strict-forbid response schema.
"""

from __future__ import annotations

import asyncio
import json
import logging
import re
from datetime import UTC, datetime
from pathlib import Path
from typing import Any, Final, Literal, cast

import orjson
from openpyxl import load_workbook
from pydantic import ValidationError

from packages.compliance.vertex_client import get_vertex_client
from packages.core.models.ratesheet import (
    ExtractionMetadata,
    NormalizedRatesheet,
    ShapeViolatingLane,
    SourceRow,
)
from packages.core.settings import Settings
from packages.ingest.prompts import (
    PROMPT_VERSION,
    SYSTEM_INSTRUCTIONS,
    USER_PROMPT_TEMPLATE,
    get_response_schema,
)

_log = logging.getLogger(__name__)

_CELL_LIMIT = 5_000
_VERTEX_BUDGET_SECONDS = 30.0
_RETRY_BACKOFF_SECONDS = 0.5
_MODEL_ID = "gemini-3.1-flash-lite"
_TEMPERATURE = 0.1

# Phase 8 §6.1 (Defect 19): canonical UN/LOCODE shape regex — identical to the
# `PortCode.code` Field pattern. Lanes whose port codes fail this regex are
# pre-scanned out of `body["lanes"]` before `NormalizedRatesheet.model_validate`
# so the Stage 2 schema gate doesn't reject the whole payload. They are routed
# to Stage 4 R1 via `shape_violating_lanes` for canonical `port_unknown_unlocode`
# rejection emission.
_UNLOCODE_SHAPE: Final = re.compile(r"^[A-Z]{2}[A-Z0-9]{3}$")


def _coerce_source_row_reference(
    raw: Any,
    *,
    sheet_name: str,
    idx: int,
) -> SourceRow:
    """Use the LLM-supplied source_row_reference when well-formed, else synthesize.

    Phase 8 §6.1 / Build Directive 1: shape-violating lanes carry their real
    Excel provenance forward so the eventual `RejectionRecord` cites the actual
    cell. Only when the LLM omitted (or malformed) the reference do we fall
    back to a positional synthesis.
    """
    if isinstance(raw, dict):
        try:
            return SourceRow.model_validate(raw)
        except ValidationError:
            pass
    # Fallback: synthesize a positional reference. `idx` is the array position
    # in body["lanes"]; +2 maps to the typical Excel layout of header row + 1.
    fallback_row = idx + 2
    return SourceRow(
        sheet_name=sheet_name[:64] if sheet_name else "Rates",
        row_number=fallback_row,
        cell_reference=f"A{fallback_row}",
    )


def _pre_scan_shape_violators(
    body: dict[str, Any],
    sheet_name: str,
) -> list[ShapeViolatingLane]:
    """Lift shape-violating lanes out of `body['lanes']` before model_validate.

    Phase 8 §6.1 (Defect 19): `PortCode`'s regex rejects shape-violators at
    `NormalizedRatesheet.model_validate(body)`, so the Phase 7 §6.1.3
    normalizer carve-out is unreachable in production. Pre-scanning here lifts
    them out cleanly; `_validate` re-injects them into Stage 4 R1 as synthetic
    `LaneRecord` carriers for canonical `port_unknown_unlocode` emission.

    Per Build Directive 1, the violator's real `source_row_reference` (emitted
    by the LLM for the originating cell) is preserved so the eventual
    `RejectionRecord` cites the true Excel coordinate, not a positional guess.
    """
    surviving: list[dict[str, Any]] = []
    violating: list[ShapeViolatingLane] = []
    for idx, lane in enumerate(body.get("lanes", [])):
        origin_code = (lane.get("origin_port") or {}).get("code", "")
        dest_code = (lane.get("destination_port") or {}).get("code", "")
        if (
            isinstance(origin_code, str)
            and isinstance(dest_code, str)
            and _UNLOCODE_SHAPE.match(origin_code)
            and _UNLOCODE_SHAPE.match(dest_code)
        ):
            surviving.append(lane)
            continue
        violating.append(
            ShapeViolatingLane(
                lane_id=str(lane.get("lane_id") or f"shape_violator_{idx}"),
                raw_origin_code=str(origin_code or ""),
                raw_destination_code=str(dest_code or ""),
                source_row_reference=_coerce_source_row_reference(
                    lane.get("source_row_reference"),
                    sheet_name=sheet_name,
                    idx=idx,
                ),
            )
        )
    body["lanes"] = surviving
    return violating


# Phase 9 §6.1 (Defect 21): canonical origin/destination header labels for the
# deterministic pre-LLM shape scan. Case-insensitive match against cell text
# after `.strip().lower()`. The primary path moves shape-violation detection
# upstream of Flash so Flash's row-omission behavior (V10 evidence) can no
# longer suppress it. When labels can't be found on a sheet, the Phase 8
# post-LLM `_pre_scan_shape_violators` carries the load as fallback.
_ORIGIN_LABELS: Final = frozenset(
    {
        "origin",
        "origin_port",
        "origin port",
        "pol",
        "load_port",
        "load port",
        "port_of_loading",
        "port of loading",
    }
)
_DESTINATION_LABELS: Final = frozenset(
    {
        "destination",
        "destination_port",
        "destination port",
        "pod",
        "discharge_port",
        "discharge port",
        "port_of_discharge",
        "port of discharge",
    }
)


def _identify_port_columns(
    cells: list[dict[str, Any]],
) -> dict[str, dict[str, int]] | None:
    """Identify origin/destination column indices per sheet by header-label match.

    Phase 9 §6.1 / Build Directive 1 (Defect 21): header detection MUST match
    on the known label set, not "the lowest string row." A banner/title cell
    above the real headers must NOT cause misidentification — either the
    labels are found on some row (correct column indices) or this function
    returns `None` and the orchestrator falls back to the Phase 8 post-LLM
    path. False negatives that look like clean passes are forbidden.

    Returns a mapping `{sheet_title: {"origin_col": int, "destination_col": int,
    "header_row": int}}` ONLY when every sheet present in the cell list has
    both an origin-label header AND a destination-label header on the SAME
    row. Returns `None` otherwise (fallback signal).
    """
    cells_by_sheet: dict[str, list[dict[str, Any]]] = {}
    for cell in cells:
        cells_by_sheet.setdefault(cell.get("sheet", ""), []).append(cell)
    if not cells_by_sheet:
        return None

    column_map: dict[str, dict[str, int]] = {}
    for sheet_title, sheet_cells in cells_by_sheet.items():
        rows: dict[int, list[dict[str, Any]]] = {}
        for cell in sheet_cells:
            row = cell.get("row")
            if isinstance(row, int):
                rows.setdefault(row, []).append(cell)
        identified: dict[str, int] | None = None
        for row_num in sorted(rows):
            origin_col: int | None = None
            dest_col: int | None = None
            for cell in rows[row_num]:
                value = cell.get("value")
                if not isinstance(value, str):
                    continue
                normalised = value.strip().lower()
                col = cell.get("col")
                if not isinstance(col, int):
                    continue
                if origin_col is None and normalised in _ORIGIN_LABELS:
                    origin_col = col
                elif dest_col is None and normalised in _DESTINATION_LABELS:
                    dest_col = col
            if origin_col is not None and dest_col is not None:
                identified = {
                    "origin_col": origin_col,
                    "destination_col": dest_col,
                    "header_row": row_num,
                }
                break
        if identified is None:
            # Per Directive 1: ANY sheet without identifiable headers fails the
            # whole workbook over to fallback. A partial-scan creates a silent
            # correctness trap.
            return None
        column_map[sheet_title] = identified
    return column_map


def _scan_cells_for_shape_violators(
    cells: list[dict[str, Any]],
    column_map: dict[str, dict[str, int]],
) -> tuple[list[dict[str, Any]], list[ShapeViolatingLane]]:
    """Lift shape-violating rows out of the cell list before Flash sees them.

    Phase 9 §6.1 / Build Directive 2 (Defect 21): every cell on a violating
    `(sheet, row)` is removed from the surviving cell list. Flash physically
    cannot re-report a row it never saw, so the post-LLM `_pre_scan_shape_violators`
    is expected to return an empty list for rows handled here — the union
    dedup is belt-and-suspenders.

    The cell list already carries true coordinates from `_build_cell_list`, so
    the lifted `ShapeViolatingLane` inherits real Excel provenance natively
    (Build Directive 1 from Phase 8 satisfied without positional synthesis).
    """
    violating_rows: dict[tuple[str, int], dict[str, dict[str, Any]]] = {}
    for cell in cells:
        sheet = cell.get("sheet")
        row = cell.get("row")
        col = cell.get("col")
        if not isinstance(sheet, str) or not isinstance(row, int) or not isinstance(col, int):
            continue
        slots = column_map.get(sheet)
        if slots is None:
            continue
        if row <= slots["header_row"]:
            continue
        role: str | None = None
        if col == slots["origin_col"]:
            role = "origin"
        elif col == slots["destination_col"]:
            role = "destination"
        if role is None:
            continue
        value = cell.get("value")
        text = str(value) if value is not None else ""
        if _UNLOCODE_SHAPE.match(text):
            continue
        slot = violating_rows.setdefault((sheet, row), {})
        slot[role] = cell

    if not violating_rows:
        return cells, []

    violators: list[ShapeViolatingLane] = []
    for (sheet, row), role_cells in sorted(violating_rows.items()):
        origin_cell = role_cells.get("origin")
        dest_cell = role_cells.get("destination")
        # The anchor coord must come from a VIOLATING cell — pin it before the
        # recovery loop fills in the non-violating side.
        anchor_cell = origin_cell or dest_cell
        # Recover whichever side wasn't a violator from the original cells so
        # `raw_*_code` reflects what was actually in the spreadsheet.
        slots = column_map[sheet]
        if origin_cell is None or dest_cell is None:
            for cell in cells:
                if cell.get("sheet") != sheet or cell.get("row") != row:
                    continue
                if origin_cell is None and cell.get("col") == slots["origin_col"]:
                    origin_cell = cell
                elif dest_cell is None and cell.get("col") == slots["destination_col"]:
                    dest_cell = cell
        assert anchor_cell is not None  # at least one violator side is present
        origin_value = str(origin_cell.get("value", "")) if origin_cell else ""
        dest_value = str(dest_cell.get("value", "")) if dest_cell else ""
        # Bound raw_*_code to ShapeViolatingLane's max_length=64 field.
        violators.append(
            ShapeViolatingLane(
                lane_id=f"shape_violator_{sheet}_{row}"[:64],
                raw_origin_code=(origin_value or " ")[:64],
                raw_destination_code=(dest_value or " ")[:64],
                source_row_reference=SourceRow(
                    sheet_name=str(sheet)[:64] or "Rates",
                    row_number=row,
                    cell_reference=str(anchor_cell.get("coord") or f"A{row}"),
                ),
            )
        )

    surviving = [
        cell for cell in cells if (cell.get("sheet"), cell.get("row")) not in violating_rows
    ]
    return surviving, violators


class ExcelTooLargeError(Exception):
    """Raised when the workbook exceeds the Phase 2 cell cap (5,000)."""


class ExtractionError(Exception):
    """Raised when Gemini Flash returned no parseable result within the budget."""


def _serialize_cell_value(value: Any) -> Any:
    """Coerce openpyxl cell values into JSON-safe primitives."""
    if value is None or isinstance(value, str | int | float | bool):
        return value
    if isinstance(value, datetime):
        return value.isoformat()
    # date, time, Decimal — coerce via str so the model sees stable text.
    return str(value)


def _build_cell_list(file_path: Path) -> list[dict[str, Any]]:
    """Walk the workbook and produce the cell list passed to Gemini.

    Enforces _CELL_LIMIT. Raises ExcelTooLargeError on overflow.
    """
    workbook = load_workbook(filename=file_path, read_only=True, data_only=True)
    try:
        cells: list[dict[str, Any]] = []
        for sheet in workbook.worksheets:
            merged_ranges = (
                {r.coord: r.coord for r in sheet.merged_cells.ranges}
                if not workbook.read_only
                else {}
            )
            for row in sheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue
                    coord = cell.coordinate
                    cells.append(
                        {
                            "sheet": sheet.title,
                            "row": cell.row,
                            "col": cell.column,
                            "coord": coord,
                            "value": _serialize_cell_value(cell.value),
                            "is_merged_anchor": coord in merged_ranges,
                            "merge_range": merged_ranges.get(coord),
                        }
                    )
                    if len(cells) > _CELL_LIMIT:
                        raise ExcelTooLargeError(
                            f"cell count exceeded {_CELL_LIMIT} (still counting at {coord})"
                        )
        return cells
    finally:
        workbook.close()


def _build_prompt(job_id: str, prospect_id: str, cells: list[dict[str, Any]]) -> str:
    """Render the user-prompt template with the JSON-serialized cell list."""
    return USER_PROMPT_TEMPLATE.format(
        job_id=job_id,
        prospect_id=prospect_id,
        prompt_version=PROMPT_VERSION,
        cells_json=orjson.dumps(cells).decode("utf-8"),
    )


async def _generate_with_retry(client: Any, contents: str, response_schema: dict[str, Any]) -> Any:
    """Call generate_content; retry once with 500ms backoff on transient 5xx errors."""
    from google.genai import types

    config = types.GenerateContentConfig(
        temperature=_TEMPERATURE,
        response_mime_type="application/json",
        response_schema=response_schema,
        system_instruction=SYSTEM_INSTRUCTIONS,
    )
    for attempt in range(2):
        try:
            return await asyncio.wait_for(
                client.aio.models.generate_content(
                    model=_MODEL_ID,
                    contents=contents,
                    config=config,
                ),
                timeout=_VERTEX_BUDGET_SECONDS,
            )
        except TimeoutError:
            _log.warning("excel_extractor: Vertex AI timeout on attempt %d", attempt + 1)
            if attempt == 0:
                await asyncio.sleep(_RETRY_BACKOFF_SECONDS)
                continue
            raise ExtractionError("Vertex AI exceeded the 30s budget") from None
        except Exception as exc:
            status = getattr(exc, "code", None) or getattr(exc, "status_code", None)
            if attempt == 0 and isinstance(status, int) and 500 <= status < 600:
                _log.warning("excel_extractor: Vertex AI %s on attempt 1; retrying once", status)
                await asyncio.sleep(_RETRY_BACKOFF_SECONDS)
                continue
            raise
    raise ExtractionError("unreachable")


def _parse_response_text(response: Any) -> dict[str, Any]:
    """Extract the JSON body from a Gemini response object."""
    text = getattr(response, "text", None)
    if text is None:
        # Fall back to inspecting candidates / parts in case the SDK surface
        # changes between minor versions.
        candidates = getattr(response, "candidates", None) or []
        for cand in candidates:
            parts = getattr(getattr(cand, "content", None), "parts", None) or []
            for part in parts:
                inline = getattr(part, "text", None)
                if inline:
                    text = inline
                    break
            if text:
                break
    if not text:
        raise ExtractionError("Vertex AI returned no text content")
    return cast(dict[str, Any], json.loads(text))


async def extract_excel_payload(
    file_path: Path,
    job_id: str,
    prospect_id: str,
    settings: Settings,
) -> tuple[NormalizedRatesheet, ExtractionMetadata]:
    """Run Stage 2 extraction on the given .xlsx file.

    Returns the parsed NormalizedRatesheet plus its ExtractionMetadata. Raises
    ExcelTooLargeError on oversized workbooks and ExtractionError on Vertex AI
    failure or schema-violating model output.
    """
    cells = await asyncio.to_thread(_build_cell_list, file_path)
    cell_count_for_metadata = len(cells)

    # Phase 9 §6.1 (Defect 21): deterministic pre-LLM shape scan on labeled-
    # column workbooks. The cell list already carries true coordinates from
    # `_build_cell_list`, so lifted violators inherit real Excel provenance
    # without positional synthesis. When `_identify_port_columns` returns
    # `None` (any sheet lacks identifiable origin/destination headers), the
    # workbook falls back to the Phase 8 post-LLM path verbatim.
    column_map = _identify_port_columns(cells)
    pre_llm_violators: list[ShapeViolatingLane] = []
    if column_map is not None:
        cells, pre_llm_violators = _scan_cells_for_shape_violators(cells, column_map)

    contents = _build_prompt(job_id, prospect_id, cells)
    client = get_vertex_client(settings)
    response = await _generate_with_retry(client, contents, get_response_schema())
    body = _parse_response_text(response)

    # Force the metadata fields to known-good values so the upstream prompt
    # cannot smuggle a different model identifier or prompt version into the
    # provenance record.
    extracted_at = datetime.now(UTC)
    metadata = ExtractionMetadata(
        extractor_model=cast(Literal["gemini-3.1-flash-lite"], _MODEL_ID),
        extracted_at=extracted_at,
        prompt_version=PROMPT_VERSION,
        cell_count=cell_count_for_metadata,
    )

    # Overwrite extractor-supplied identifiers with the authoritative ones.
    body["job_id"] = job_id
    body["prospect_id"] = prospect_id
    body["extraction_metadata"] = metadata.model_dump(mode="json")
    body["schema_version"] = "onramp.v1"

    # Phase 8 §6.1 (Defect 19): post-LLM pre-scan, retained as fallback after
    # Phase 9 moved the primary path upstream of Flash. When the pre-LLM scan
    # ran (column_map was not None), Flash never saw the violating rows so
    # this call is expected to return an empty list — the union/dedup below
    # is belt-and-suspenders per Build Directive 2. When the pre-LLM path
    # punted (column_map was None), this is the sole detector.
    sheet_label = "Rates"
    post_llm_violators = _pre_scan_shape_violators(body, sheet_label)

    # Phase 9 §6.1: union the two violator sources, deduped by
    # (sheet_name, row_number) from source_row_reference — robust against
    # lane_id divergence between pre-LLM synthetic IDs and any Flash-supplied
    # IDs the post-LLM path may carry.
    seen_rows: set[tuple[str, int]] = set()
    shape_violating: list[ShapeViolatingLane] = []
    for v in (*pre_llm_violators, *post_llm_violators):
        key = (v.source_row_reference.sheet_name, v.source_row_reference.row_number)
        if key in seen_rows:
            continue
        seen_rows.add(key)
        shape_violating.append(v)

    # Phase 7 §6.1.1 (Defect 18a): force-overwrite, not setdefault. The Stage 2
    # LLM occasionally smuggles non-empty values into these fields with
    # invented rule_ids like "INVALID_PORT_CODE". Stage 4 is the sole source
    # of rejections; conformal/flagged are populated downstream.
    body["conformal_scores"] = {}
    body["flagged_for_review"] = []
    body["deterministically_rejected"] = []
    body["shape_violating_lanes"] = [v.model_dump(mode="json") for v in shape_violating]

    # Phase 8 §6.2 (Defect 20): wrap model_validate so any Pydantic violation
    # surfaces as ExtractionError. Without this, ValidationError propagates
    # past `_extract`'s `except (ExcelTooLargeError, ExtractionError)` handler
    # and the job sticks at `status=extracting` indefinitely.
    try:
        payload = NormalizedRatesheet.model_validate(body)
    except ValidationError as exc:
        raise ExtractionError(
            f"stage2_schema_violation: {exc.error_count()} pydantic error(s)"
        ) from exc
    return payload, metadata
